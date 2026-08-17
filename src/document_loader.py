"""Document loading + chunking for the RAG pipeline.

Supports markdown (with optional YAML frontmatter), plain text, PDF and DOCX.
Used by both:
  - `ingest.py` (batch ingest of `data/<role>/*.md`)
  - `auth/upload UI` (live ingest of files uploaded by Admin / HR)
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import yaml

from .config import CHUNK_OVERLAP, CHUNK_SIZE, DATA_DIR

FRONTMATTER_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n(.*)$", re.DOTALL)

SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}


@dataclass
class Chunk:
    chunk_id: str
    text: str
    metadata: dict = field(default_factory=dict)


# --- Frontmatter / metadata ----------------------------------------------


def _parse_frontmatter(raw: str, file_path: Path) -> tuple[dict, str]:
    """Split YAML frontmatter from body, falling back to filename defaults."""
    match = FRONTMATTER_RE.match(raw)
    if not match:
        return (
            {
                "title": file_path.stem.replace("_", " ").title(),
                "role": file_path.parent.name,
                "system": "both",
                "source_id": file_path.stem.upper(),
            },
            raw.strip(),
        )

    meta = yaml.safe_load(match.group(1)) or {}
    body = match.group(2).strip()
    meta.setdefault("title", file_path.stem.replace("_", " ").title())
    meta.setdefault("role", file_path.parent.name)
    meta.setdefault("system", "both")
    meta.setdefault("source_id", file_path.stem.upper())
    return meta, body


# --- Chunking -------------------------------------------------------------


def _chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware sliding window chunker."""
    if not text:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    buffer = ""

    for para in paragraphs:
        candidate = f"{buffer}\n\n{para}".strip() if buffer else para
        if len(candidate) <= size:
            buffer = candidate
            continue

        if buffer:
            chunks.append(buffer)
            tail = buffer[-overlap:] if overlap and len(buffer) > overlap else ""
            buffer = f"{tail}\n\n{para}".strip() if tail else para
        else:
            for i in range(0, len(para), size - overlap):
                chunks.append(para[i : i + size])
            buffer = ""

    if buffer:
        chunks.append(buffer)
    return chunks


# --- Per-format text extractors ------------------------------------------


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # local import keeps cold-start fast

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - some PDFs raise on weird fonts
            pages.append("")
    return "\n\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document  # python-docx

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


def _extract_text_for_suffix(suffix: str, data: bytes) -> str:
    s = suffix.lower()
    if s in {".md", ".markdown", ".txt"}:
        return data.decode("utf-8", errors="replace")
    if s == ".pdf":
        return _extract_pdf(data)
    if s == ".docx":
        return _extract_docx(data)
    raise ValueError(f"Unsupported file type: {suffix}")


# --- Public chunk builders -----------------------------------------------


def _build_chunks(body: str, meta: dict) -> list[Chunk]:
    """Chunk `body` and stamp each chunk with the same metadata."""
    pieces = _chunk_text(body)
    if not pieces:
        return []
    out: list[Chunk] = []
    for idx, piece in enumerate(pieces):
        meta_out = {
            "title": str(meta.get("title", "Untitled")),
            "role": str(meta.get("role", "common")).lower(),
            "system": str(meta.get("system", "both")).lower(),
            "source_id": str(meta.get("source_id", "DOC")).upper(),
            "source_path": str(meta.get("source_path", "")),
            "uploaded_by": str(meta.get("uploaded_by", "")),
            "chunk_index": idx,
            "total_chunks": len(pieces),
        }
        chunk_id = f"{meta_out['source_id']}-{idx:04d}"
        out.append(Chunk(chunk_id=chunk_id, text=piece, metadata=meta_out))
    return out


def chunks_from_file_bytes(
    filename: str,
    data: bytes,
    *,
    role: str,
    system: str = "both",
    title: str | None = None,
    source_id: str | None = None,
    uploaded_by: str | None = None,
    source_path: str | None = None,
) -> list[Chunk]:
    """Build chunks from in-memory bytes (used by the upload UI)."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Unsupported file type '{suffix}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    body = _extract_text_for_suffix(suffix, data).strip()
    if not body:
        return []

    stem = Path(filename).stem
    meta = {
        "title": title or stem.replace("_", " ").replace("-", " ").title(),
        "role": role,
        "system": system,
        "source_id": (source_id or stem).upper(),
        "source_path": source_path or filename,
        "uploaded_by": uploaded_by or "",
    }
    return _build_chunks(body, meta)


def load_documents(data_dir: Path = DATA_DIR) -> Iterable[Chunk]:
    """Walk `data_dir` and yield Chunks for every supported file."""
    if not data_dir.exists():
        return

    for path in sorted(data_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue

        data = path.read_bytes()
        suffix = path.suffix.lower()

        if suffix in {".md", ".markdown", ".txt"}:
            raw = data.decode("utf-8", errors="replace")
            meta, body = _parse_frontmatter(raw, path)
        else:
            body = _extract_text_for_suffix(suffix, data)
            meta = {
                "title": path.stem.replace("_", " ").title(),
                "role": path.parent.name,
                "system": "both",
                "source_id": path.stem.upper(),
            }

        meta["source_path"] = str(path.relative_to(data_dir.parent))
        for chunk in _build_chunks(body, meta):
            yield chunk
